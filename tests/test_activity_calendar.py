import io
import unittest
from datetime import date, time

from openpyxl import load_workbook

from activity_calendar import (
    ACTIVITY_LIST_HEADERS,
    LOGGED,
    MUNICIPALITY_NOT_RECORDED,
    PLANNED,
    STATUS_LEGEND,
    build_calendar_entries,
    calendar_filename,
    calendar_summary_counts,
    calendar_xlsx_bytes,
    group_by_day,
    logged_calendar_rows,
    municipality_ward_compact,
    municipality_ward_label,
    planned_calendar_rows,
)
from activity_records import CONFIRMED_DUPLICATE
from leadership_reporting import entry_date
from week_dates import month_bounds, next_month_key, previous_month_key


MONTH_KEY = "2026-09"
MONTH_START, MONTH_END = month_bounds(MONTH_KEY)


def entry(**overrides):
    doc = {
        "id": "1", "person_id": "willem-p", "name": "Willem P",
        "activity_date": "2026-09-16", "ward": "Ward 7", "venue": "Bezville Hall",
        "type": "Door to Door", "type_display": "Door to Door",
        "start_time": "09:00", "end_time": "11:00", "campaign_id": "",
    }
    doc.update(overrides)
    return doc


def campaign(**overrides):
    doc = {
        "id": "c1", "name": "Soup Kitchen Drive", "person_id": "spokazi-m",
        "municipality": "Amahlathi", "wards": ["Ward 14"], "submission_status": "submitted",
        "planned_activities": [],
    }
    doc.update(overrides)
    return doc


class MunicipalityWardLabelTests(unittest.TestCase):
    def test_combines_municipality_and_single_ward(self):
        self.assertEqual(municipality_ward_label("Raymond Mhlaba", "Ward 7"), "Raymond Mhlaba Ward 7")

    def test_combines_municipality_and_multiple_wards(self):
        self.assertEqual(
            municipality_ward_label("Amahlathi", ["Ward 2", "Ward 14"]),
            "Amahlathi Ward 2, Amahlathi Ward 14",
        )

    def test_raymond_mhlaba_ward_7_and_amahlathi_ward_7_remain_distinct(self):
        a = municipality_ward_label("Raymond Mhlaba", "Ward 7")
        b = municipality_ward_label("Amahlathi", "Ward 7")
        self.assertNotEqual(a, b)
        self.assertEqual({a, b}, {"Raymond Mhlaba Ward 7", "Amahlathi Ward 7"})

    def test_unknown_municipality_shows_explicit_not_recorded_text(self):
        self.assertEqual(municipality_ward_label("", "Ward 7"), MUNICIPALITY_NOT_RECORDED)
        self.assertEqual(municipality_ward_label(None, None), MUNICIPALITY_NOT_RECORDED)

    def test_municipality_known_no_ward_shows_municipality_alone(self):
        self.assertEqual(municipality_ward_label("Amahlathi", ""), "Amahlathi")

    def test_never_double_prefixes_a_legacy_municipality_only_ward(self):
        self.assertEqual(municipality_ward_label("Amahlathi", "Amahlathi"), "Amahlathi")


class LoggedCalendarRowsTests(unittest.TestCase):
    def test_uses_real_stored_activity_data(self):
        rows = logged_calendar_rows(
            [entry()], {}, {"willem-p": "Raymond Mhlaba"}, MONTH_START, MONTH_END, entry_date,
        )
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["date"], date(2026, 9, 16))
        self.assertEqual(row["activity"], "Door to Door")
        self.assertEqual(row["venue"], "Bezville Hall")
        self.assertEqual(row["candidate"], "Willem P")
        self.assertEqual(row["status"], LOGGED)
        self.assertEqual(row["municipality_ward"], "Raymond Mhlaba Ward 7")

    def test_status_is_logged_not_planned(self):
        rows = logged_calendar_rows([entry()], {}, {}, MONTH_START, MONTH_END, entry_date)
        self.assertEqual(rows[0]["status"], "LOGGED")
        self.assertNotEqual(rows[0]["status"], "PLANNED")

    def test_excludes_activities_outside_the_month(self):
        rows = logged_calendar_rows(
            [entry(activity_date="2026-08-31"), entry(activity_date="2026-10-01")],
            {}, {}, MONTH_START, MONTH_END, entry_date,
        )
        self.assertEqual(rows, [])

    def test_confirmed_duplicate_excluded_using_shared_rule(self):
        rows = logged_calendar_rows(
            [entry(duplicate_review_status=CONFIRMED_DUPLICATE)],
            {}, {}, MONTH_START, MONTH_END, entry_date,
        )
        self.assertEqual(rows, [], "a confirmed duplicate must never appear on the calendar")

    def test_missing_time_is_never_invented(self):
        rows = logged_calendar_rows(
            [entry(start_time=None, end_time=None)], {}, {}, MONTH_START, MONTH_END, entry_date,
        )
        self.assertEqual(rows[0]["time_label"], "Time not recorded")
        self.assertEqual(rows[0]["start_time"], "")

    def test_municipality_not_inferred_from_venue_or_name(self):
        rows = logged_calendar_rows(
            [entry(venue="Amahlathi Community Hall", name="Someone in Raymond Mhlaba")],
            {}, {},  # no roster municipality known for this person_id
            MONTH_START, MONTH_END, entry_date,
        )
        self.assertEqual(rows[0]["municipality"], "")
        self.assertEqual(rows[0]["municipality_ward"], MUNICIPALITY_NOT_RECORDED)

    def test_campaign_name_resolved_by_campaign_id(self):
        rows = logged_calendar_rows(
            [entry(campaign_id="c1")], {"c1": "Soup Kitchen Drive"}, {}, MONTH_START, MONTH_END, entry_date,
        )
        self.assertEqual(rows[0]["campaign_name"], "Soup Kitchen Drive")


class PlannedCalendarRowsTests(unittest.TestCase):
    def test_uses_planned_activity_row_data(self):
        c = campaign(planned_activities=[
            {"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": "On the street corner"},
        ])
        rows = planned_calendar_rows([c], {"spokazi-m": "Spokazi M"}, MONTH_START, MONTH_END)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["date"], date(2026, 9, 20))
        self.assertEqual(row["activity"], "Soup Kitchen")
        self.assertEqual(row["venue"], "On the street corner")
        self.assertEqual(row["candidate"], "Spokazi M")
        self.assertEqual(row["campaign_name"], "Soup Kitchen Drive")
        self.assertEqual(row["status"], PLANNED)
        self.assertEqual(row["municipality_ward"], "Amahlathi Ward 14")

    def test_status_is_planned_not_logged(self):
        c = campaign(planned_activities=[{"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""}])
        rows = planned_calendar_rows([c], {}, MONTH_START, MONTH_END)
        self.assertEqual(rows[0]["status"], "PLANNED")

    def test_planned_row_never_marked_as_completed_reporting(self):
        # A planned row must never be classifiable as LOGGED/reportable via
        # the shared duplicate/reportable rule — it isn't even an `entries`
        # document, so it can never enter is_reportable_activity's domain.
        c = campaign(planned_activities=[{"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""}])
        rows = planned_calendar_rows([c], {}, MONTH_START, MONTH_END)
        self.assertNotIn("duplicate_review_status", rows[0])
        self.assertEqual(rows[0]["status"], PLANNED)

    def test_draft_campaigns_excluded(self):
        c = campaign(submission_status="draft", planned_activities=[
            {"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""},
        ])
        rows = planned_calendar_rows([c], {}, MONTH_START, MONTH_END)
        self.assertEqual(rows, [])

    def test_activities_outside_month_excluded(self):
        c = campaign(planned_activities=[
            {"id": "p1", "date": "2026-08-31", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""},
            {"id": "p2", "date": "2026-10-01", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""},
        ])
        rows = planned_calendar_rows([c], {}, MONTH_START, MONTH_END)
        self.assertEqual(rows, [])

    def test_missing_planned_time_is_never_invented(self):
        c = campaign(planned_activities=[{"id": "p1", "date": "2026-09-20", "time": "", "activity_type": "Soup Kitchen", "area": ""}])
        rows = planned_calendar_rows([c], {}, MONTH_START, MONTH_END)
        self.assertEqual(rows[0]["time_label"], "Time not recorded")

    def test_no_campaign_municipality_shows_not_recorded(self):
        c = campaign(municipality="", planned_activities=[
            {"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""},
        ])
        rows = planned_calendar_rows([c], {}, MONTH_START, MONTH_END)
        self.assertEqual(rows[0]["municipality_ward"], MUNICIPALITY_NOT_RECORDED)


class BuildCalendarEntriesTests(unittest.TestCase):
    def test_combines_logged_and_planned_without_deduping(self):
        # No reliable per-row link exists between a planned_activities item
        # and a logged entries document anywhere in this app, so both a
        # planned row and an unrelated logged row for the same day must
        # both appear — never silently collapsed.
        e = entry(activity_date="2026-09-20", type_display="Soup Kitchen")
        c = campaign(planned_activities=[
            {"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""},
        ])
        rows = build_calendar_entries([e], [c], {}, {}, {}, MONTH_START, MONTH_END, entry_date)
        self.assertEqual(len(rows), 2)
        self.assertEqual({r["status"] for r in rows}, {"LOGGED", "PLANNED"})

    def test_status_filter_planned_only(self):
        e = entry()
        c = campaign(planned_activities=[{"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""}])
        rows = build_calendar_entries([e], [c], {}, {}, {}, MONTH_START, MONTH_END, entry_date, status_filter="PLANNED")
        self.assertTrue(rows)
        self.assertTrue(all(r["status"] == "PLANNED" for r in rows))

    def test_status_filter_logged_only(self):
        e = entry()
        c = campaign(planned_activities=[{"id": "p1", "date": "2026-09-20", "time": "09:00", "activity_type": "Soup Kitchen", "area": ""}])
        rows = build_calendar_entries([e], [c], {}, {}, {}, MONTH_START, MONTH_END, entry_date, status_filter="logged")
        self.assertTrue(rows)
        self.assertTrue(all(r["status"] == "LOGGED" for r in rows))

    def test_sorted_chronologically(self):
        e1 = entry(activity_date="2026-09-20", start_time="14:00")
        e2 = entry(activity_date="2026-09-05", start_time="09:00")
        rows = build_calendar_entries([e1, e2], [], {}, {}, {}, MONTH_START, MONTH_END, entry_date)
        self.assertEqual([r["date"] for r in rows], [date(2026, 9, 5), date(2026, 9, 20)])


class GroupByDayTests(unittest.TestCase):
    def test_groups_multiple_entries_on_the_same_day(self):
        rows = [
            {"date": date(2026, 9, 20), "start_time": "09:00"},
            {"date": date(2026, 9, 20), "start_time": "14:00"},
            {"date": date(2026, 9, 21), "start_time": "09:00"},
        ]
        grouped = group_by_day(rows)
        self.assertEqual(len(grouped[date(2026, 9, 20)]), 2)
        self.assertEqual(len(grouped[date(2026, 9, 21)]), 1)


class MonthNavigationIntegrationTests(unittest.TestCase):
    def test_next_and_previous_month_key_round_trip(self):
        self.assertEqual(previous_month_key(next_month_key(MONTH_KEY)), MONTH_KEY)


class CalendarXlsxTests(unittest.TestCase):
    def _rows(self):
        e = entry(activity_date="2026-09-16")
        c = campaign(planned_activities=[
            {"id": "p1", "date": "2026-09-20", "time": "10:00", "activity_type": "Soup Kitchen", "area": "On the street corner"},
        ])
        return build_calendar_entries(
            [e], [c], {}, {"willem-p": "Raymond Mhlaba"}, {"spokazi-m": "Spokazi M"},
            MONTH_START, MONTH_END, entry_date,
        )

    def test_produces_a_genuine_xlsx_with_two_worksheets(self):
        payload = calendar_xlsx_bytes(self._rows(), MONTH_KEY)
        self.assertTrue(payload.startswith(b"PK\x03\x04"), "must be a genuine .xlsx, not CSV/text")
        wb = load_workbook(io.BytesIO(payload))
        self.assertEqual(wb.sheetnames, ["Calendar", "Activity List"])

    def test_calendar_sheet_shows_clean_header_block(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Calendar"]
        self.assertEqual(ws["A2"].value, "Democratic Alliance")
        self.assertEqual(ws["A3"].value, "Ntsikana Constituency")
        self.assertEqual(ws["A4"].value, "Activity Calendar")
        self.assertEqual(ws["A5"].value, "September 2026")

    def test_calendar_header_is_not_oversized(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Calendar"]
        # The weekday header (and grid) must start within the first ~9 rows
        # — a handful of compact header lines, not a sprawling banner.
        self.assertEqual(ws["A9"].value, "Monday")

    def test_month_summary_reconciles_exactly_with_the_rows(self):
        rows = self._rows()
        total, logged, planned = calendar_summary_counts(rows)
        self.assertEqual((total, logged, planned), (2, 1, 1))
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(rows, MONTH_KEY)))
        ws = wb["Calendar"]
        self.assertEqual(ws["A6"].value, f"Total Activities: {total}  |  Logged: {logged}  |  Planned: {planned}")

    def test_empty_month_shows_plain_message_not_zero_counts(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes([], MONTH_KEY)))
        ws = wb["Calendar"]
        self.assertEqual(ws["A6"].value, "No activities scheduled for this month.")

    def test_legend_present_and_print_safe(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Calendar"]
        self.assertEqual(ws["A7"].value, STATUS_LEGEND)
        self.assertIn("✓", STATUS_LEGEND)
        self.assertIn("○", STATUS_LEGEND)
        self.assertIn("Logged", STATUS_LEGEND)
        self.assertIn("Planned", STATUS_LEGEND)

    def test_calendar_cell_text_uses_compact_symbol_and_ward_form(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Calendar"]
        day16 = next(c for row in ws.iter_rows(min_row=10) for c in row if c.value and str(c.value).startswith("16\n"))
        self.assertIn("✓ 09:00 Door to Door · Raymond Mhlaba W7", str(day16.value))
        self.assertNotIn("(LOGGED)", str(day16.value))
        self.assertNotIn("Ward 7", str(day16.value), "the compact grid must abbreviate to W7, not spell out Ward 7")

    def test_missing_time_never_invented_and_stays_compact(self):
        docs_row = {
            "date": date(2026, 9, 21), "start_time": "", "end_time": "", "time_label": "Time not recorded",
            "activity": "Info Table", "municipality": "Amahlathi", "ward": "Ward 4",
            "municipality_ward": "Amahlathi Ward 4", "venue": "", "candidate": "", "campaign_name": "", "status": PLANNED,
        }
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes([docs_row], MONTH_KEY)))
        ws = wb["Calendar"]
        day21 = next(c for row in ws.iter_rows(min_row=10) for c in row if c.value and str(c.value).startswith("21\n"))
        self.assertIn("○ Info Table · Amahlathi W4", str(day21.value))
        self.assertNotIn("Time not recorded", str(day21.value), "the Calendar grid must never repeat 'Time not recorded'")

    def test_busy_day_row_grows_and_never_drops_activities(self):
        busy_docs = [
            entry(activity_date="2026-09-10", start_time=f"{9+i:02d}:00", end_time=f"{10+i:02d}:00", type_display=f"Activity {i}")
            for i in range(8)
        ]
        rows = logged_calendar_rows(busy_docs, {}, {"willem-p": "Raymond Mhlaba"}, MONTH_START, MONTH_END, entry_date)
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(rows, MONTH_KEY)))
        ws = wb["Calendar"]
        day10 = next(c for row in ws.iter_rows(min_row=10) for c in row if c.value and str(c.value).startswith("10\n"))
        for i in range(8):
            self.assertIn(f"Activity {i}", str(day10.value), "every activity on a busy day must still be present, never dropped")
        row_height = ws.row_dimensions[day10.row].height
        self.assertGreater(row_height, 80, "a busy day's row must grow taller than the default")

    def test_days_outside_the_month_are_shaded_not_prominent(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Calendar"]
        # 1 Sep 2026 is a Tuesday: cell A10 (Monday of the first grid row) is
        # outside the month and must be blank, not a real day-1 cell.
        self.assertFalse(ws["A10"].value)
        self.assertEqual(ws["A10"].fill.patternType, "solid")

    def test_activity_list_headers_exact(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Activity List"]
        self.assertEqual([c.value for c in ws[1]], ACTIVITY_LIST_HEADERS)

    def test_activity_list_date_and_time_cells_are_real_types_with_sa_format(self):
        import datetime as dt
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Activity List"]
        rows_by_status = {row[9].value: row for row in ws.iter_rows(min_row=2)}
        logged_row = rows_by_status["LOGGED"]
        self.assertIsInstance(logged_row[0].value, dt.date)
        self.assertEqual(logged_row[0].number_format, "dd/mm/yyyy")
        self.assertIsInstance(logged_row[1].value, dt.time)
        self.assertEqual(logged_row[1].number_format, "HH:MM")
        self.assertEqual(logged_row[3].value, "Raymond Mhlaba")  # MUNICIPALITY column
        self.assertEqual(logged_row[4].value, "Ward 7")  # WARD column, separate

    def test_planned_row_has_no_end_time_and_status_planned(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        ws = wb["Activity List"]
        planned_row = next(row for row in ws.iter_rows(min_row=2) if row[9].value == "PLANNED")
        self.assertIsNone(planned_row[2].value)  # TIME END
        self.assertEqual(planned_row[3].value, "Amahlathi")
        self.assertEqual(planned_row[4].value, "Ward 14")

    def test_no_formulas_outside_the_header_block(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")

    def test_only_the_six_header_lines_are_merged_on_the_calendar_sheet(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        merged = sorted(str(r) for r in wb["Calendar"].merged_cells.ranges)
        self.assertEqual(merged, [f"A{n}:G{n}" for n in range(2, 8)])
        self.assertEqual(list(wb["Activity List"].merged_cells.ranges), [])

    def test_logo_embedded_when_available(self):
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(self._rows(), MONTH_KEY)))
        # static/logo.png exists in this repo (the same file the Coordinator/
        # Leadership weekly report workbook already embeds).
        self.assertEqual(len(wb["Calendar"]._images), 1)

    def test_activity_list_can_be_omitted(self):
        payload = calendar_xlsx_bytes(self._rows(), MONTH_KEY, include_activity_list=False)
        wb = load_workbook(io.BytesIO(payload))
        self.assertEqual(wb.sheetnames, ["Calendar"])

    def test_activity_list_rows_alternate_shading(self):
        rows = self._rows() + self._rows()  # ensure at least 3 data rows
        wb = load_workbook(io.BytesIO(calendar_xlsx_bytes(rows, MONTH_KEY)))
        ws = wb["Activity List"]
        first_data_row_fill = ws.cell(row=2, column=1).fill.patternType
        second_data_row_fill = ws.cell(row=3, column=1).fill.patternType
        self.assertNotEqual(first_data_row_fill, second_data_row_fill, "alternating rows must actually differ")


class CalendarFilenameTests(unittest.TestCase):
    def test_filename_reflects_selected_month(self):
        self.assertEqual(calendar_filename("2026-09"), "Ntsikana_Activity_Calendar_September_2026.xlsx")
        self.assertEqual(calendar_filename("2026-10"), "Ntsikana_Activity_Calendar_October_2026.xlsx")

    def test_filename_never_says_current(self):
        self.assertNotIn("current", calendar_filename("2026-09").lower())


class MunicipalityWardCompactTests(unittest.TestCase):
    def test_abbreviates_ward_number(self):
        self.assertEqual(municipality_ward_compact("Raymond Mhlaba Ward 7"), "Raymond Mhlaba W7")

    def test_abbreviates_every_ward_in_a_multi_ward_list(self):
        self.assertEqual(
            municipality_ward_compact("Amahlathi Ward 2, Amahlathi Ward 14"),
            "Amahlathi W2, Amahlathi W14",
        )

    def test_leaves_municipality_not_recorded_untouched(self):
        self.assertEqual(municipality_ward_compact(MUNICIPALITY_NOT_RECORDED), MUNICIPALITY_NOT_RECORDED)


if __name__ == "__main__":
    unittest.main()
