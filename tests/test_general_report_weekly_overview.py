import asyncio
import copy
import io
import os
import unittest
from types import SimpleNamespace

try:
    import fastapi  # noqa: F401
    from bson import ObjectId

    HAS_API_DEPS = True
except ModuleNotFoundError:
    ObjectId = None
    HAS_API_DEPS = False

try:
    import openpyxl

    HAS_OPENPYXL = True
except ModuleNotFoundError:
    HAS_OPENPYXL = False


# Phase 5.1: "Weekly Overview" sheet + trend chart on the general admin
# workbook (/api/admin/export.xlsx) ONLY — this file never touches, imports
# from, or asserts against smartsheet_reporting.py; that workbook/format is
# a separate, explicitly out-of-scope hard boundary for this phase.
@unittest.skipUnless(HAS_API_DEPS and HAS_OPENPYXL, "FastAPI/openpyxl dependencies are not installed")
class WeeklyOverviewSheetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("MONGO_URI", "mongodb://127.0.0.1:1")
        os.environ.setdefault("ADMIN_PIN", "1234")
        os.environ.setdefault("JWT_SECRET", "local-test-secret")
        global appmod
        import main as appmod

    def setUp(self):
        self.original_entries_col = appmod.entries_col
        self.original_roster_col = appmod.roster_col
        self.entries = FakeCollection()
        self.roster = FakeCollection()
        self.roster.docs = [
            {"_id": ObjectId(), "name": "Test Candidate", "ward": "Ward 1", "name_slug": "test-candidate"},
        ]
        appmod.entries_col = self.entries
        appmod.roster_col = self.roster

    def tearDown(self):
        appmod.entries_col = self.original_entries_col
        appmod.roster_col = self.original_roster_col

    def _seed(self, week_key, count, campaign_id=None, day="mon"):
        for _ in range(count):
            self.entries.docs.append({
                "_id": ObjectId(), "person_id": "test-candidate", "name": "Test Candidate",
                "ward": "Ward 1", "day": day, "type": "Door to Door", "type_display": "Door to Door",
                "notes": None, "week_key": week_key, "week_label": "x", "activity_date": week_key,
                "venue": "Hall", "campaign_id": campaign_id,
                "submitted_at": "2026-09-01T10:00:00+00:00",
            })

    async def _export(self, week_key=None):
        resp = await appmod.admin_export_xlsx(week_key=week_key, _=True)
        chunks = [c async for c in resp.body_iterator]
        return openpyxl.load_workbook(io.BytesIO(b"".join(chunks)))

    def test_weekly_overview_sheet_exists_alongside_report(self):
        self._seed("2026-08-30", 2)
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        self.assertEqual(wb.sheetnames, ["Report", "Weekly Overview"])

    def test_report_sheet_content_unaffected(self):
        self._seed("2026-08-30", 2)
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        self.assertEqual(
            wb["Report"]["A2"].value, "Ntsikana Constituency - Weekly Ward Activity Report",
        )

    def test_weekly_overview_chronological_order_and_totals(self):
        self._seed("2026-08-16", 3)
        self._seed("2026-08-23", 5)
        self._seed("2026-08-30", 2)
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        ws = wb["Weekly Overview"]
        rows = [tuple(r) for r in ws.iter_rows(min_row=1, max_row=4, values_only=True)]
        self.assertEqual(rows[0], ("Week", "Activities", "Change vs Previous Week"))
        weeks = [r[0] for r in rows[1:]]
        self.assertEqual(weeks, sorted(weeks), "weeks must be shown chronologically")
        self.assertEqual([r[1] for r in rows[1:]], [3, 5, 2])

    def test_weekly_overview_counts_campaign_and_ordinary_activities_together(self):
        self._seed("2026-08-30", 2, campaign_id=None)
        self._seed("2026-08-30", 3, campaign_id="camp1")
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        ws = wb["Weekly Overview"]
        self.assertEqual(ws["B2"].value, 5)

    def test_weekly_overview_excludes_weeks_after_the_reports_own_week(self):
        self._seed("2026-08-30", 2)
        self._seed("2026-09-06", 9)  # after the reported week — must not leak in
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        ws = wb["Weekly Overview"]
        weeks_shown = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value]
        self.assertEqual(len(weeks_shown), 1)
        self.assertEqual(ws["B2"].value, 2)

    def test_weekly_overview_change_vs_previous_week(self):
        self._seed("2026-08-16", 3)
        self._seed("2026-08-23", 5)
        self._seed("2026-08-30", 2)
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        ws = wb["Weekly Overview"]
        self.assertEqual(ws["C2"].value, "-", "first shown week has no prior point")
        self.assertEqual(ws["C3"].value, 2)
        self.assertEqual(ws["C4"].value, -3)

    def test_weekly_overview_has_a_chart_referencing_the_activities_column(self):
        self._seed("2026-08-16", 3)
        self._seed("2026-08-23", 5)
        wb = asyncio.run(self._export(week_key="2026-08-23"))
        ws = wb["Weekly Overview"]
        self.assertEqual(len(ws._charts), 1)
        chart = ws._charts[0]
        series = chart.series[0]
        # The series' value reference must point at column B (Activities) on
        # this same sheet — not an unrelated/empty range.
        self.assertIn("Weekly Overview", series.val.numRef.f)
        self.assertIn("$B$", series.val.numRef.f)

    def test_weekly_overview_with_no_data_up_to_report_week_has_no_chart(self):
        wb = asyncio.run(self._export(week_key="2026-08-30"))
        ws = wb["Weekly Overview"]
        self.assertEqual(len(ws._charts), 0)
        self.assertEqual(ws["A1"].value, "Week")

    def test_export_never_touches_smartsheet_reporting_functions(self):
        # Hard boundary: the general admin workbook must never call into the
        # SmartSheet workbook builders — those are a separate, frozen format.
        smartsheet_fns = [
            "smartsheet_xlsx_bytes",
            "smartsheet_workbook_all_categories_bytes",
            "smartsheet_csv_bytes",
        ]
        originals = {name: getattr(appmod, name) for name in smartsheet_fns}
        calls = []
        try:
            for name in smartsheet_fns:
                setattr(appmod, name, lambda *a, _n=name, **k: calls.append(_n))
            self._seed("2026-08-30", 2)
            asyncio.run(self._export(week_key="2026-08-30"))
        finally:
            for name, fn in originals.items():
                setattr(appmod, name, fn)
        self.assertEqual(calls, [])


class AsyncCursor:
    def __init__(self, docs):
        self.docs = [copy.deepcopy(doc) for doc in docs]

    def __aiter__(self):
        self._iter = iter(self.docs)
        return self

    async def __anext__(self):
        try:
            return next(self._iter)
        except StopIteration:
            raise StopAsyncIteration


class FakeCollection:
    def __init__(self):
        self.docs = []

    async def create_index(self, *args, **kwargs):
        return "idx"

    def find(self, query=None, projection=None):
        query = query or {}
        docs = [doc for doc in self.docs if matches(doc, query)]
        return AsyncCursor(docs)

    async def find_one(self, query):
        for doc in self.docs:
            if matches(doc, query):
                return copy.deepcopy(doc)
        return None

    async def insert_one(self, doc):
        stored = copy.deepcopy(doc)
        stored["_id"] = ObjectId()
        self.docs.append(stored)
        return SimpleNamespace(inserted_id=stored["_id"])


def matches(doc, query):
    return all(doc.get(key) == value for key, value in query.items())


if __name__ == "__main__":
    unittest.main()
